#!/usr/bin/env python3
"""把一份真实资料变成文本：PDF / Word / Excel / 图片 OCR。

为什么要有它：老师手里的资料是 PDF、Word、Excel、手机拍的卷子——标准库解不了；
让模型每次现写脚本既慢又容易写错。常见的活这里一次干完，
干不了的（版式太怪、扫描质量太差）如实说，让模型自己写脚本接手。

用法（stdout 一律是 JSON，便于上层处理）：

    extract.py probe   <文件>              # 这是什么？要不要 OCR？
    extract.py text    <文件> [--max N]    # 直接读文本类文件
    extract.py pdf     <文件> [--max N] [--ocr]
    extract.py docx    <文件> [--max N]
    extract.py xlsx    <文件> [--max N]
    extract.py ocr     <文件> [--lang chi_sim+eng] [--max N]

输出：{ok, kind, pages?, chars, truncated, needs_ocr, text, notes: [...]}
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".jsonl", ".json", ".log", ".tex"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
# 一页少于这么多字符，基本可以断定是扫描件（只有页码或没有文字层）
SCANNED_PAGE_CHARS = 20
DEFAULT_MAX = 20000


def note(notes: list[str], message: str) -> None:
    notes.append(message)


def read_text_file(path: Path, notes: list[str]) -> str:
    raw = path.read_bytes()
    try:
        import chardet  # type: ignore

        guess = chardet.detect(raw[:200_000])
        encoding = guess.get("encoding") or "utf-8"
        confidence = guess.get("confidence") or 0
        if encoding.lower() in {"gb2312", "gbk", "gb18030"}:
            encoding = "gb18030"
        text = raw.decode(encoding, errors="replace")
        if encoding.lower() not in {"utf-8", "ascii"}:
            note(notes, f"文件不是 UTF-8，按 {encoding} 读的（置信度 {confidence:.2f}）")
        return text
    except ImportError:
        note(notes, "没装 chardet，按 UTF-8 读（乱码就装一下）")
        return raw.decode("utf-8", errors="replace")


def pdf_text(path: Path, notes: list[str]) -> tuple[str, int, int]:
    """返回 (文本, 页数, 空页数)"""
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        raise RuntimeError("没装 pdfplumber：请在仓库根跑 pnpm venv") from None

    chunks: list[str] = []
    empty = 0
    with pdfplumber.open(str(path)) as pdf:
        pages = len(pdf.pages)
        for index, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if len(text.strip()) < SCANNED_PAGE_CHARS:
                empty += 1
            chunks.append(f"（第 {index} 页）\n{text}")
        if empty == pages and pages > 0:
            note(notes, "整份 PDF 都没有文字层，多半是扫描件：用 --ocr")
    return "\n\n".join(chunks), pages, empty


def tessdata_env() -> dict[str, str]:
    """仓库自带的语言包（data/tessdata）优先——系统包装不了就要 sudo，这里不用。"""
    env = dict(os.environ)
    if env.get("TESSDATA_PREFIX"):
        return env
    here = Path(__file__).resolve().parent.parent
    for candidate in (here / "data" / "tessdata", Path.cwd() / "data" / "tessdata"):
        if candidate.is_dir():
            env["TESSDATA_PREFIX"] = str(candidate)
            break
    return env


def tesseract_available() -> tuple[bool, list[str]]:
    binary = shutil.which("tesseract")
    if binary is None:
        return False, []
    try:
        out = subprocess.run(
            [binary, "--list-langs"], capture_output=True, text=True, timeout=30, env=tessdata_env()
        )
        langs = [line.strip() for line in out.stdout.splitlines()[1:] if line.strip()]
        return True, langs
    except (OSError, subprocess.SubprocessError):
        return False, []


def pick_langs(wanted: str, notes: list[str]) -> str:
    ok, langs = tesseract_available()
    if not ok:
        raise RuntimeError("系统里没有 tesseract：装一下（apt install tesseract-ocr）")
    requested = [part for part in wanted.split("+") if part]
    usable = [part for part in requested if part in langs]
    missing = [part for part in requested if part not in langs]
    if missing:
        note(
            notes,
            f"缺少语言包 {'、'.join(missing)}（现在只能认 {'、'.join(usable) or '无'}）；"
            "中文要装 tesseract-ocr-chi-sim / tesseract-ocr-chi-tra",
        )
    if not usable:
        # 退到现有语言里最合适的一个，并说清退化到什么
        fallback = "eng" if "eng" in langs else (langs[0] if langs else "")
        if fallback == "":
            raise RuntimeError("tesseract 没有任何可用语言包")
        note(notes, f"改用 {fallback} 识别")
        return fallback
    return "+".join(usable)


def ocr_image(path: Path, lang: str, notes: list[str]) -> str:
    language = pick_langs(lang, notes)
    result = subprocess.run(
        ["tesseract", str(path), "-", "-l", language],
        capture_output=True,
        text=True,
        timeout=300,
        env=tessdata_env(),
    )
    if result.returncode != 0:
        raise RuntimeError(f"tesseract 失败：{result.stderr.strip()[:200]}")
    return result.stdout


def ocr_pdf(path: Path, lang: str, notes: list[str], max_pages: int = 20) -> str:
    """扫描件：逐页渲染成图再 OCR（pypdfium2 是 pdfplumber 的依赖，通常已经在）"""
    try:
        import pypdfium2 as pdfium  # type: ignore
    except ImportError:
        raise RuntimeError("没装 pypdfium2：请在仓库根跑 pnpm venv") from None

    language = pick_langs(lang, notes)
    doc = pdfium.PdfDocument(str(path))
    pieces: list[str] = []
    pages = len(doc)
    if pages > max_pages:
        note(notes, f"只 OCR 了前 {max_pages} 页（共 {pages} 页），后面的自己翻")
    for index in range(min(pages, max_pages)):
        page = doc[index]
        image = page.render(scale=2.0).to_pil()
        tmp = path.with_suffix(f".p{index + 1}.png")
        image.save(tmp)
        try:
            text = subprocess.run(
                ["tesseract", str(tmp), "-", "-l", language],
                capture_output=True,
                text=True,
                timeout=300,
                env=tessdata_env(),
            ).stdout
        finally:
            tmp.unlink(missing_ok=True)
        pieces.append(f"（第 {index + 1} 页）\n{text}")
    return "\n\n".join(pieces)


def docx_text(path: Path, notes: list[str]) -> str:
    try:
        import docx  # type: ignore
    except ImportError:
        raise RuntimeError("没装 python-docx：请在仓库根跑 pnpm venv") from None

    document = docx.Document(str(path))
    pieces = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip() != ""]
    for index, table in enumerate(document.tables, start=1):
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        pieces.append(f"（表 {index}）\n" + "\n".join(rows))
    if not pieces:
        note(notes, "文档里没有段落文字（可能整篇是图片：转成 PDF 再试 OCR）")
    return "\n\n".join(pieces)


def xlsx_text(path: Path, notes: list[str]) -> str:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        raise RuntimeError("没装 openpyxl：请在仓库根跑 pnpm venv") from None

    book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    pieces: list[str] = []
    for sheet in book.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value) for value in row]
            if any(cell.strip() != "" for cell in cells):
                rows.append("\t".join(cells))
        pieces.append(f"（工作表 {sheet.title}）\n" + "\n".join(rows))
    book.close()
    if len(pieces) == 0:
        note(notes, "工作簿是空的")
    return "\n\n".join(pieces)


def probe(path: Path) -> dict:
    notes: list[str] = []
    suffix = path.suffix.lower()
    info: dict = {
        "kind": "unknown",
        "pages": None,
        "needs_ocr": False,
        "notes": notes,
    }
    if suffix == ".pdf":
        info["kind"] = "pdf"
        try:
            _text, pages, empty = pdf_text(path, notes)
            info["pages"] = pages
            if pages > 0 and empty == pages:
                info["needs_ocr"] = True
                note(notes, "没有文字层，需要 OCR")
        except (RuntimeError, OSError) as error:
            note(notes, str(error))
    elif suffix == ".docx":
        info["kind"] = "docx"
    elif suffix in {".xlsx", ".xlsm"}:
        info["kind"] = "xlsx"
    elif suffix in IMAGE_SUFFIXES:
        info["kind"] = "image"
        info["needs_ocr"] = True
        ok, langs = tesseract_available()
        note(notes, f"OCR 可用语言：{'、'.join(langs) if ok else '没有 tesseract'}")
    elif suffix in TEXT_SUFFIXES:
        info["kind"] = "text"
    else:
        note(notes, f"不认识的后缀 {suffix or '（没有后缀）'}；先按文本试试，不行就自己写脚本")
    return info


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(json.dumps({"ok": False, "error": "用法：extract.py <probe|text|pdf|docx|xlsx|ocr> <文件>"}, ensure_ascii=False))
        return 2

    mode = argv[1]
    rest = argv[2:]
    max_chars = DEFAULT_MAX
    lang = "chi_sim+eng"
    use_ocr = False
    paths: list[str] = []
    index = 0
    while index < len(rest):
        token = rest[index]
        if token == "--max":
            index += 1
            max_chars = int(rest[index]) if index < len(rest) else max_chars
        elif token == "--lang":
            index += 1
            lang = rest[index] if index < len(rest) else lang
        elif token == "--ocr":
            use_ocr = True
        else:
            paths.append(token)
        index += 1

    if not paths:
        print(json.dumps({"ok": False, "error": "没给文件"}, ensure_ascii=False))
        return 2

    path = Path(paths[0]).expanduser()
    if not path.is_file():
        print(json.dumps({"ok": False, "error": f"读不到文件：{path}"}, ensure_ascii=False))
        return 1

    notes: list[str] = []
    try:
        if mode == "probe":
            result = probe(path)
            result.update({"ok": True, "path": str(path), "bytes": path.stat().st_size})
            print(json.dumps(result, ensure_ascii=False))
            return 0

        kind = mode
        pages = None
        text = ""
        if mode == "text":
            text = read_text_file(path, notes)
        elif mode == "pdf":
            text, pages, empty = pdf_text(path, notes)
            if use_ocr and pages is not None and empty > 0:
                notes.append("有页面没有文字层，改用 OCR 读整份")
                text = ocr_pdf(path, lang, notes)
        elif mode == "docx":
            text = docx_text(path, notes)
        elif mode == "xlsx":
            text = xlsx_text(path, notes)
        elif mode == "ocr":
            text = ocr_image(path, lang, notes) if path.suffix.lower() in IMAGE_SUFFIXES else ocr_pdf(path, lang, notes)
        else:
            print(json.dumps({"ok": False, "error": f"不认识的模式：{mode}"}, ensure_ascii=False))
            return 2

        total = len(text)
        truncated = total > max_chars
        print(
            json.dumps(
                {
                    "ok": True,
                    "kind": kind,
                    "path": str(path),
                    "bytes": path.stat().st_size,
                    "pages": pages,
                    "chars": total,
                    "truncated": truncated,
                    "needs_ocr": False,
                    "text": text[:max_chars],
                    "notes": notes,
                },
                ensure_ascii=False,
            )
        )
        return 0
    except (RuntimeError, OSError, ValueError) as error:
        print(
            json.dumps(
                {"ok": False, "error": str(error), "notes": notes, "path": str(path)},
                ensure_ascii=False,
            )
        )
        return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
